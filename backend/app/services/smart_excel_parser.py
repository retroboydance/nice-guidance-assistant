"""
Smart Excel and CSV parser for NICE Baseline Assessment Tools (BATs).
Capable of parsing raw, unmodified NICE spreadsheets as downloaded directly
from the NICE website, as well as CSV exports.
"""

from dataclasses import dataclass, field
import io
import re
from typing import BinaryIO, Union
import openpyxl


@dataclass
class ParsedRecommendation:
    rec_number: str
    rec_text: str
    rec_year: str = ""
    section_title: str = ""


@dataclass
class ParsedGuidelineData:
    title: str
    code: str
    recommendations: list[ParsedRecommendation] = field(default_factory=list)


class SmartExcelParser:
    """Intelligently detects sheets, headers, and recommendations from raw NICE BAT files."""

    REC_NUM_REGEX = re.compile(r"^\d+(?:\.\d+)+$")
    CODE_REGEX = re.compile(r"(?:\b|_)(NG|CG|QS|PH|MTG|DG|HST|TA)\s*(\d+)(?:\b|_)", re.IGNORECASE)

    @classmethod
    def parse(
        cls,
        file_input: Union[bytes, BinaryIO, str],
        filename: str = "",
    ) -> ParsedGuidelineData:
        """Parse raw bytes, file-like object, or file path."""
        if isinstance(file_input, str) and file_input.endswith(".csv"):
            with open(file_input, "rb") as f:
                return cls._parse_csv(f.read(), filename or file_input)
        elif isinstance(file_input, (bytes, bytearray)):
            if filename.endswith(".csv") or (file_input[:100].find(b",") != -1 and b"PK\x03\x04" not in file_input[:4]):
                return cls._parse_csv(file_input, filename)
            return cls._parse_xlsx(io.BytesIO(file_input), filename)
        elif hasattr(file_input, "read"):
            data = file_input.read()
            if filename.endswith(".csv") or b"PK\x03\x04" not in data[:4]:
                return cls._parse_csv(data, filename)
            return cls._parse_xlsx(io.BytesIO(data), filename)
        else:
            with open(file_input, "rb") as f:
                return cls._parse_xlsx(io.BytesIO(f.read()), filename or str(file_input))

    @classmethod
    def _parse_xlsx(cls, stream: io.BytesIO, filename: str = "") -> ParsedGuidelineData:
        wb = openpyxl.load_workbook(stream, data_only=True, read_only=True)

        sheet = cls._select_worksheet(wb)
        rows = list(sheet.iter_rows(values_only=True))

        if not rows:
            raise ValueError("Spreadsheet contains no data.")

        title = cls._extract_title(rows, filename)
        code = cls._extract_code(title, filename)

        header_row_idx, col_map = cls._detect_columns(rows)

        recommendations: list[ParsedRecommendation] = []
        current_section = ""

        rec_col = col_map["rec_num"]
        text_col = col_map["rec_text"]
        year_col = col_map.get("rec_year")

        for row in rows[header_row_idx + 1:]:
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue

            rec_num_raw = row[rec_col] if rec_col < len(row) and row[rec_col] is not None else ""
            rec_text_raw = row[text_col] if text_col < len(row) and row[text_col] is not None else ""
            rec_year_raw = row[year_col] if year_col and year_col < len(row) and row[year_col] is not None else ""

            rec_num = str(rec_num_raw).strip()
            rec_text = cls._clean_text(str(rec_text_raw))
            rec_year = str(rec_year_raw).strip()

            if not rec_num and rec_text:
                if re.match(r"^\d+(\.\d+)?\s+[A-Za-z]", rec_text):
                    current_section = rec_text
                continue

            if cls.REC_NUM_REGEX.match(rec_num) and rec_text:
                recommendations.append(
                    ParsedRecommendation(
                        rec_number=rec_num,
                        rec_text=rec_text,
                        rec_year=rec_year,
                        section_title=current_section,
                    )
                )

        if not recommendations:
            raise ValueError("No valid NICE recommendations found in the spreadsheet.")

        return ParsedGuidelineData(
            title=title,
            code=code,
            recommendations=recommendations,
        )

    @classmethod
    def _parse_csv(cls, raw_bytes: bytes, filename: str = "") -> ParsedGuidelineData:
        import csv

        text = ""
        for encoding in ["utf-8-sig", "utf-16", "cp1252", "latin-1"]:
            try:
                candidate = raw_bytes.decode(encoding)
                if "\x00" not in candidate:
                    text = candidate
                    break
            except UnicodeDecodeError:
                continue

        if not text:
            text = raw_bytes.decode("utf-8", errors="ignore").replace("\x00", "")

        reader = csv.reader(io.StringIO(text))
        rows = [row for row in reader if any(str(cell).strip() for cell in row)]

        if not rows:
            raise ValueError("CSV contains no data.")

        title = cls._extract_title(rows, filename)
        code = cls._extract_code(title, filename)
        header_row_idx, col_map = cls._detect_columns(rows)

        recommendations: list[ParsedRecommendation] = []
        current_section = ""

        rec_col = col_map["rec_num"]
        text_col = col_map["rec_text"]
        year_col = col_map.get("rec_year")

        for row in rows[header_row_idx + 1:]:
            rec_num_raw = row[rec_col] if rec_col < len(row) else ""
            rec_text_raw = row[text_col] if text_col < len(row) else ""
            rec_year_raw = row[year_col] if year_col and year_col < len(row) else ""

            rec_num = str(rec_num_raw).strip()
            rec_text = cls._clean_text(str(rec_text_raw))
            rec_year = str(rec_year_raw).strip()

            if not rec_num and rec_text and re.match(r"^\d+(\.\d+)?\s+[A-Za-z]", rec_text):
                current_section = rec_text
                continue

            if cls.REC_NUM_REGEX.match(rec_num) and rec_text:
                recommendations.append(
                    ParsedRecommendation(
                        rec_number=rec_num,
                        rec_text=rec_text,
                        rec_year=rec_year,
                        section_title=current_section,
                    )
                )

        return ParsedGuidelineData(
            title=title,
            code=code,
            recommendations=recommendations,
        )

    @classmethod
    def _select_worksheet(cls, wb: openpyxl.Workbook):
        """Find the relevant Data Sheet among NICE workbook sheets."""
        for name in wb.sheetnames:
            lower = name.lower()
            if "data" in lower or "baseline" in lower or "assessment" in lower:
                return wb[name]

        for name in wb.sheetnames:
            lower = name.lower()
            if not any(skip in lower for skip in ["intro", "dropdown", "total", "summary", "table 1"]):
                return wb[name]

        return wb.active

    @classmethod
    def _extract_title(cls, rows: list, filename: str) -> str:
        """Extract guideline title from row 1 or filename."""
        if rows and rows[0]:
            first_cell = str(rows[0][0] or "").strip()
            if "baseline assessment tool for" in first_cell.lower():
                return re.sub(r"^baseline assessment tool for\s+", "", first_cell, flags=re.IGNORECASE).strip()
            elif len(first_cell) > 5 and not cls.REC_NUM_REGEX.match(first_cell):
                return first_cell

        if filename:
            cleaned = re.sub(r"\.(xlsx|xls|csv)$", "", filename, flags=re.IGNORECASE)
            return cleaned.replace("-", " ").replace("_", " ").title()

        return "NICE Guideline"

    @classmethod
    def _extract_code(cls, title: str, filename: str) -> str:
        """Extract NICE code such as NG253, CG191, QS213."""
        match = cls.CODE_REGEX.search(title)
        if match:
            return f"{match.group(1).upper()}{match.group(2)}"

        if filename:
            match = cls.CODE_REGEX.search(filename)
            if match:
                return f"{match.group(1).upper()}{match.group(2)}"

        return ""

    @classmethod
    def _detect_columns(cls, rows: list) -> tuple[int, dict[str, int]]:
        """Identify which columns contain rec_num, rec_text, and rec_year."""
        for idx, row in enumerate(rows[:10]):
            col_map: dict[str, int] = {}
            for col_idx, cell in enumerate(row):
                if cell is None:
                    continue
                val = str(cell).strip().lower()

                if any(k in val for k in ["guideline recommendation", "recommendation number", "rec number", "guideline number", "rec_num", "recommendation no", "rec no", "ref"]):
                    col_map["rec_num"] = col_idx
                elif "rec_text" not in col_map and any(k in val for k in ["nice recommendation", "recommendation text", "guideline text", "rec_text", "recommendation details"]):
                    col_map["rec_text"] = col_idx
                elif "rec_year" not in col_map and any(k in val for k in ["year of recommendation", "publish info", "rec_year", "year"]):
                    col_map["rec_year"] = col_idx

            if "rec_text" not in col_map:
                for col_idx, cell in enumerate(row):
                    if cell is None:
                        continue
                    val = str(cell).strip().lower()
                    if val in ["recommendation", "recommendations"] and col_idx != col_map.get("rec_num"):
                        col_map["rec_text"] = col_idx
                        break

            if "rec_num" in col_map and "rec_text" in col_map:
                return idx, col_map

        return 1, {"rec_text": 0, "rec_num": 1, "rec_year": 2}

    @staticmethod
    def _clean_text(text: str) -> str:
        """Clean non-breaking spaces and formatting artefacts."""
        if not text:
            return ""
        cleaned = text.replace("\xa0", " ").replace("\r\n", "\n").replace("\r", "\n")
        cleaned = cleaned.replace("\ufffd", "• ")
        return cleaned.strip()
